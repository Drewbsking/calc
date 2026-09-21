r"""Plot seasonal factors from one Excel workbook as yearly PNG heatmaps.

Install once: py -m pip install openpyxl matplotlib
Run from Windows Terminal: py .\seasonal_factor_plots.py
Reads the workbook locally; does not modify it or replace existing output files.
"""

# =====================================================
# CHANGE THE SETTINGS IN THIS BOX.
# EXCEL_FILE_PATH: paste the full path to your input workbook.
# Keep the r and quotation marks.
# =====================================================
EXCEL_FILE_PATH = r"C:\CHANGE_ME\Seasonal Factor Table.xlsx"
SHEET_NAME = "Seasonal_Factor_Table"
START_YEAR = 2021
END_YEAR = 2023
OUTPUT_FOLDER = ""  # Blank: create seasonal_heatmaps beside the workbook.
DPI = 300
# Example: EXCEL_FILE_PATH = r"C:\Work\Seasonal Factor Table.xlsx"
# No code changes are needed below this line.

import math
import os
from pathlib import Path


MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
DAY_COLUMNS = ("MONFAC", "TUEFAC", "WEDFAC", "THURFAC", "FRIFAC", "SATFAC", "SUNFAC")
DAY_LABELS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")


def check_years(start_year, end_year):
    if (type(start_year) is not int or type(end_year) is not int
            or not 1 <= start_year <= end_year <= 9999):
        raise ValueError("START_YEAR and END_YEAR must be whole years in increasing order (1 to 9999).")


def reject_links(path):
    for part in (path, *path.parents):
        if part.is_symlink() or (hasattr(part, "is_junction") and part.is_junction()):
            raise ValueError(f"Symbolic links and directory junctions are not processed: {part}")


def prepare_factors(rows, start_year, end_year):
    """Validate worksheet rows and return 12 x 7 matrices and one shared scale."""
    check_years(start_year, end_year)
    rows = iter(rows)
    header = [str(value).strip() if value is not None else "" for value in next(rows, ())]
    required = ("YEAR", "MONTH", *DAY_COLUMNS)
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError("Missing worksheet columns: " + ", ".join(missing))
    if any(header.count(name) != 1 for name in required):
        raise ValueError("Required worksheet column names must be unique.")
    columns = {name: header.index(name) for name in required}
    matrices = {}
    seen = set()
    values = []
    for row_number, row in enumerate(rows, start=2):
        if all(value is None or value == "" for value in row):
            continue
        row = tuple(row) + (None,) * max(0, len(header) - len(row))
        raw_year = row[columns["YEAR"]]
        try:
            year = float(raw_year)
        except (TypeError, ValueError):
            year = math.nan
        if isinstance(raw_year, bool) or not math.isfinite(year) or not year.is_integer():
            raise ValueError(f"Row {row_number}: YEAR must be a whole year.")
        year = int(year)
        if not start_year <= year <= end_year:
            continue
        month = str(row[columns["MONTH"]]).strip().title()
        if month not in MONTHS:
            raise ValueError(f"Row {row_number}: MONTH must use Jan through Dec.")
        if (year, month) in seen:
            raise ValueError(f"Duplicate month: {month} {year}. Keep one row per year and month.")
        seen.add((year, month))
        matrix = matrices.setdefault(year, [[math.nan] * 7 for _ in MONTHS])
        for day, name in enumerate(DAY_COLUMNS):
            raw_value = row[columns[name]]
            if raw_value is None or (isinstance(raw_value, str) and not raw_value.strip()):
                continue
            try:
                value = float(raw_value)
            except (TypeError, ValueError):
                value = math.nan
            if isinstance(raw_value, bool) or not math.isfinite(value) or value <= 0:
                raise ValueError(f"Row {row_number}, {name}: factor must be a positive, finite number or blank.")
            matrix[MONTHS.index(month)][day] = value
            values.append(value)
    empty_years = [year for year in range(start_year, end_year + 1)
                   if year not in matrices or not any(math.isfinite(value)
                   for row in matrices[year] for value in row)]
    if empty_years:
        raise ValueError("No numeric factors for requested year(s): " + ", ".join(map(str, empty_years)))
    # Symmetric bounds keep 1.0 neutral and every year's colors comparable.
    spread = max(abs(min(values) - 1), abs(max(values) - 1), 0.01)
    return {year: matrices[year] for year in sorted(matrices)}, (1 - spread, 1 + spread)


def load_factors(source, sheet_name, start_year, end_year):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Missing Excel package. Install with: py -m pip install openpyxl matplotlib") from error
    try:
        with Path(source).open("rb") as source_file:
            workbook = load_workbook(source_file, read_only=True, data_only=True)
            try:
                if sheet_name not in workbook.sheetnames:
                    raise ValueError(f"Sheet {sheet_name!r} not found. Available sheets: {', '.join(workbook.sheetnames)}")
                return prepare_factors(workbook[sheet_name].values, start_year, end_year)
            finally:
                workbook.close()
    except (OSError, ValueError):
        raise
    except Exception as error:
        raise ValueError(f"Could not read workbook: {error}") from error


def make_figure(year, matrix, bounds):
    """Use an offscreen canvas so saving does not open or wait for plot windows."""
    try:
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        from matplotlib.colors import LinearSegmentedColormap, Normalize
        from matplotlib.figure import Figure
    except ImportError as error:
        raise RuntimeError("Missing plotting package. Install with: py -m pip install openpyxl matplotlib") from error
    figure = Figure(figsize=(10, 7), layout="constrained")
    FigureCanvasAgg(figure)
    axes = figure.subplots()
    cmap = LinearSegmentedColormap.from_list("seasonal", ("#397795", "#f7f7f7", "#bf4c52"))
    cmap.set_bad("#e5e7eb")
    norm = Normalize(*bounds)
    plot = axes.imshow(matrix, cmap=cmap, norm=norm, aspect="auto", interpolation="nearest")
    axes.set_xticks(range(7), labels=DAY_LABELS)
    axes.set_yticks(range(12), labels=MONTHS)
    axes.set_xticks([i - 0.5 for i in range(8)], minor=True)
    axes.set_yticks([i - 0.5 for i in range(13)], minor=True)
    axes.grid(which="minor", color="white", linewidth=0.8)
    axes.tick_params(which="both", length=0)
    for month, row in enumerate(matrix):
        for day, value in enumerate(row):
            present = math.isfinite(value)
            red, green, blue, _ = cmap(norm(value)) if present else (1, 1, 1, 1)
            color = "white" if 0.2126 * red + 0.7152 * green + 0.0722 * blue < 0.5 else "#172033"
            axes.text(day, month, f"{value:.2f}" if present else "—", ha="center", va="center",
                      fontsize=10, color=color, fontweight="bold" if present and 0.95 <= value <= 1.05 else "normal")
    axes.set_title(f"Seasonal Factors Heatmap ({year}) — Common Scale", pad=14)
    axes.set_xlabel("Day of the Week")
    axes.set_ylabel("Month")
    figure.colorbar(plot, ax=axes, label="Factor (centered at 1.00)", pad=0.025)
    figure.supxlabel("Bold: factors from 0.95 to 1.05  •  Gray / —: no data", fontsize=10)
    return figure


def generate_heatmaps(excel_path, sheet_name=SHEET_NAME, start_year=START_YEAR,
                      end_year=END_YEAR, output_folder=OUTPUT_FOLDER, dpi=DPI):
    if not str(excel_path).strip() or "CHANGE_ME" in str(excel_path):
        raise ValueError("Set EXCEL_FILE_PATH in the settings box to your input workbook.")
    check_years(start_year, end_year)
    if type(dpi) is not int or not 72 <= dpi <= 1200:
        raise ValueError("DPI must be a whole number from 72 to 1200.")
    source = Path(excel_path).absolute()
    reject_links(source)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError(f"Excel workbook not found or not an .xlsx file: {source}")
    destination = Path(output_folder).absolute() if output_folder else source.parent / "seasonal_heatmaps"
    reject_links(destination)
    matrices, bounds = load_factors(source, sheet_name, start_year, end_year)
    outputs = [destination / f"heatmap_{year}_common_scale.png" for year in matrices]
    for path in outputs:
        if os.path.lexists(path):
            raise FileExistsError(f"Output already exists: {path}. Choose another OUTPUT_FOLDER.")
    destination.mkdir(parents=True, exist_ok=True)
    for (year, matrix), path in zip(matrices.items(), outputs):
        figure = make_figure(year, matrix, bounds)
        try:
            with path.open("xb") as output:
                try:
                    figure.savefig(output, format="png", dpi=dpi)
                    output.flush()
                except BaseException:
                    output.close()
                    path.unlink()
                    raise
        finally:
            figure.clear()
        print(f"Saved: {path}")
    print(f"Summary: {len(outputs)} heatmap(s) saved with a shared scale centered at 1.00.")
    return outputs


def main():
    try:
        generate_heatmaps(EXCEL_FILE_PATH, SHEET_NAME, START_YEAR, END_YEAR, OUTPUT_FOLDER, DPI)
    except (OSError, ValueError, RuntimeError) as error:
        print(f"Error: {error}")
        return 1
    except KeyboardInterrupt:
        print("Cancelled. Completed heatmaps are kept; any incomplete image is removed.")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
